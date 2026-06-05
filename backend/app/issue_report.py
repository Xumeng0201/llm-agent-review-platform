from __future__ import annotations

import html
from io import BytesIO

from docx import Document

from .docx_helpers import apply_global_yahei_font, format_now_beijing
from .models import ReviewIssue, ReviewTask


def severity_label(severity: str) -> str:
    return {
        "serious": "严重问题",
        "major": "一般问题",
        "minor": "轻微问题",
    }.get(severity, severity)


def status_label(status: str) -> str:
    return {
        "open": "待处理",
        "accepted": "已采纳",
        "dismissed": "已驳回",
        "revised": "已修订",
    }.get(status, status)


def _grouped_issues(issues: list[ReviewIssue]) -> list[tuple[str, list[ReviewIssue]]]:
    return [
        ("严重问题", [x for x in issues if x.severity == "serious"]),
        ("一般问题", [x for x in issues if x.severity == "major"]),
        ("轻微问题", [x for x in issues if x.severity == "minor"]),
    ]


def build_issue_report_html(task: ReviewTask, issues: list[ReviewIssue]) -> str:
    serious = [x for x in issues if x.severity == "serious"]
    major = [x for x in issues if x.severity == "major"]
    minor = [x for x in issues if x.severity == "minor"]
    grouped_blocks: list[str] = []
    running_index = 1
    for group_title, group_issues in _grouped_issues(issues):
        if not group_issues:
            continue
        issue_blocks: list[str] = []
        for issue in group_issues:
            evidence_html = "".join(
                (
                    "<li>"
                    f"<strong>{html.escape(ev.file_name)}</strong>"
                    f"{' · 第 ' + str(ev.page_from) + ' 页' if ev.page_from is not None else ''}"
                    f"{' · ' + html.escape(ev.section_title) if ev.section_title else ''}"
                    f"<div class='quote'>{html.escape(ev.quote_text)}</div>"
                    "</li>"
                )
                for ev in issue.evidences
            )
            issue_blocks.append(
                "<div class='issue'>"
                f"<div class='issue-head'><span class='sev sev-{html.escape(issue.severity)}'>{html.escape(severity_label(issue.severity))}</span>"
                f"<span class='dim'>{html.escape(issue.dimension_title)}</span>"
                f"<span class='status'>{html.escape(status_label(issue.status))}</span></div>"
                f"<h3>{running_index}. {html.escape(issue.title)}</h3>"
                f"<p><strong>问题描述：</strong>{html.escape(issue.description)}</p>"
                f"<p><strong>形成原因：</strong>{html.escape(issue.reason or '—')}</p>"
                f"<p><strong>整改建议：</strong>{html.escape(issue.suggestion or '—')}</p>"
                f"<p><strong>补充材料：</strong>{'是' if issue.needs_supplement else '否'}　<strong>人工重点复核：</strong>{'是' if issue.manual_review else '否'}</p>"
                f"<div><strong>证据引用：</strong><ul>{evidence_html or '<li>无</li>'}</ul></div>"
                "</div>"
            )
            running_index += 1
        grouped_blocks.append(f"<h3 class='group-title'>{html.escape(group_title)}</h3>{''.join(issue_blocks)}")
    missing_materials = (task.missing_materials or "（无）").replace("\n", "<br/>")
    overall = task.overall_assessment or "（尚未生成总体判断）"
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8"/>
  <meta name="viewport" content="width=device-width, initial-scale=1"/>
  <title>{html.escape(task.name)} - 审查意见书</title>
  <style>
    body {{ font-family: ui-sans-serif, system-ui, 'Segoe UI', 'Noto Sans SC', sans-serif; background:#f8fafc; color:#0f172a; margin:0; padding:32px; }}
    .sheet {{ max-width: 980px; margin:0 auto; background:#fff; border:1px solid #e2e8f0; border-radius:16px; padding:32px; }}
    h1 {{ margin:0 0 8px; font-size:28px; }}
    .meta {{ color:#64748b; font-size:14px; margin-bottom:24px; }}
    .summary {{ display:grid; grid-template-columns: repeat(3, minmax(0,1fr)); gap:12px; margin:18px 0 24px; }}
    .card {{ padding:14px 16px; border-radius:14px; border:1px solid #e2e8f0; background:#f8fafc; }}
    .num {{ font-size:26px; font-weight:700; margin-top:6px; }}
    .issue {{ border-top:1px solid #e2e8f0; padding-top:18px; margin-top:18px; }}
    .group-title {{ margin-top:28px; margin-bottom:4px; font-size:22px; }}
    .issue-head {{ display:flex; gap:8px; flex-wrap:wrap; align-items:center; margin-bottom:10px; }}
    .sev, .dim, .status {{ display:inline-block; font-size:12px; padding:4px 10px; border-radius:999px; }}
    .sev-serious {{ background:#fee2e2; color:#991b1b; }}
    .sev-major {{ background:#fef3c7; color:#92400e; }}
    .sev-minor {{ background:#dbeafe; color:#1d4ed8; }}
    .dim {{ background:#eef2ff; color:#4338ca; }}
    .status {{ background:#e2e8f0; color:#334155; }}
    .quote {{ margin-top:6px; white-space:pre-wrap; background:#f8fafc; border:1px solid #e2e8f0; border-radius:10px; padding:10px 12px; color:#334155; }}
    ul {{ padding-left:20px; }}
  </style>
</head>
<body>
  <div class="sheet">
    <h1>{html.escape(task.name)}</h1>
    <div class="meta">项目方案智能审查意见书 · 生成时间（北京时间）：{html.escape(format_now_beijing())}</div>
    <h2>总体判断</h2>
    <div class="card" style="white-space:pre-wrap">{html.escape(overall)}</div>
    <div class="summary">
      <div class="card"><div>严重问题</div><div class="num">{len(serious)}</div></div>
      <div class="card"><div>一般问题</div><div class="num">{len(major)}</div></div>
      <div class="card"><div>轻微问题</div><div class="num">{len(minor)}</div></div>
    </div>
    <h2>建议补充材料</h2>
    <div class="card">{missing_materials}</div>
    <h2>问题清单</h2>
    {''.join(grouped_blocks) if grouped_blocks else "<p>当前暂无问题记录。</p>"}
  </div>
</body>
</html>"""


def build_issue_report_docx_bytes(task: ReviewTask, issues: list[ReviewIssue]) -> bytes:
    doc = Document()
    apply_global_yahei_font(doc, body_pt=11)

    doc.add_heading(task.name, level=0)
    ts = format_now_beijing()
    doc.add_paragraph(f"项目方案智能审查意见书 · 生成时间（北京时间）：{ts}")

    doc.add_heading("总体判断", level=1)
    doc.add_paragraph(task.overall_assessment or "（尚未生成总体判断）")

    doc.add_heading("问题统计", level=1)
    doc.add_paragraph(f"严重问题：{sum(1 for x in issues if x.severity == 'serious')}")
    doc.add_paragraph(f"一般问题：{sum(1 for x in issues if x.severity == 'major')}")
    doc.add_paragraph(f"轻微问题：{sum(1 for x in issues if x.severity == 'minor')}")

    doc.add_heading("建议补充材料", level=1)
    missing = (task.missing_materials or "").strip()
    if missing:
        for line in [x.strip() for x in missing.splitlines() if x.strip()]:
            doc.add_paragraph(line, style="List Bullet")
    else:
        doc.add_paragraph("（无）")

    doc.add_heading("问题清单", level=1)
    if not issues:
        doc.add_paragraph("当前暂无问题记录。")
    running_index = 1
    for group_title, group_issues in _grouped_issues(issues):
        if not group_issues:
            continue
        doc.add_heading(group_title, level=2)
        for issue in group_issues:
            doc.add_heading(f"{running_index}. {issue.title}", level=3)
            doc.add_paragraph(f"严重程度：{severity_label(issue.severity)}")
            doc.add_paragraph(f"所属维度：{issue.dimension_title}")
            doc.add_paragraph(f"处理状态：{status_label(issue.status)}")
            doc.add_paragraph(f"问题描述：{issue.description}")
            doc.add_paragraph(f"形成原因：{issue.reason or '—'}")
            doc.add_paragraph(f"整改建议：{issue.suggestion or '—'}")
            doc.add_paragraph(f"需补充材料：{'是' if issue.needs_supplement else '否'}；人工重点复核：{'是' if issue.manual_review else '否'}")
            doc.add_paragraph("证据引用：")
            if issue.evidences:
                for ev in issue.evidences:
                    head = ev.file_name
                    if ev.page_from is not None:
                        head += f" 第 {ev.page_from} 页"
                    if ev.section_title:
                        head += f" · {ev.section_title}"
                    doc.add_paragraph(head, style="List Bullet")
                    doc.add_paragraph(ev.quote_text)
            else:
                doc.add_paragraph("无", style="List Bullet")
            running_index += 1

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
