from __future__ import annotations

import json
import shutil
from dataclasses import dataclass
from pathlib import Path


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".htm",
    ".log",
}
IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
}


@dataclass
class ParsedDocument:
    text: str
    page_count: int | None = None
    outline_json: str | None = None
    pages: list[str] | None = None
    parser_engine: str = "builtin"
    parser_mode: str | None = None


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_table_lines(doc) -> list[str]:
    """表格内文字（许多方案材料正文写在表里，仅读 paragraphs 会得到空串）。"""
    out: list[str] = []
    try:
        for table in doc.tables:
            for row in table.rows:
                cells: list[str] = []
                for cell in row.cells:
                    cell_text = "\n".join(
                        (p.text or "").strip() for p in cell.paragraphs if (p.text or "").strip()
                    ).strip()
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    out.append(" | ".join(cells))
    except Exception:
        pass
    return out


def _read_docx_wt_fallback(path: Path) -> str:
    """从 word/document.xml 抽取所有 w:t，覆盖文本框等 python-docx 段落 API 漏掉的情况。"""
    import zipfile
    import xml.etree.ElementTree as ET

    try:
        with zipfile.ZipFile(path) as zf:
            raw = zf.read("word/document.xml")
    except Exception:
        return ""
    try:
        root = ET.fromstring(raw)
    except Exception:
        return ""
    parts: list[str] = []
    for el in root.iter(_W_NS + "t"):
        if el.text:
            parts.append(el.text)
        if el.tail:
            parts.append(el.tail)
    return " ".join(parts).strip()


def _read_docx_file(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        return ""
    try:
        doc = Document(str(path))
    except Exception:
        return ""
    lines: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style_name = ""
        try:
            style_name = (p.style.name or "").lower()
        except Exception:
            style_name = ""
        # 尽量保留 Word 中的标题语义，方便后续 chunk 切分和模板章节映射。
        if "heading" in style_name or "标题" in style_name:
            if lines and lines[-1] != "":
                lines.append("")
            lines.append(text)
            lines.append("")
        else:
            lines.append(text)
    table_lines = _docx_table_lines(doc)
    if table_lines:
        if lines and lines[-1] != "":
            lines.append("")
        lines.append("【表格内容】")
        lines.extend(table_lines)
    merged = "\n".join(lines).strip()
    if merged:
        return merged
    return _read_docx_wt_fallback(path)


def _read_pdf_file(path: Path) -> tuple[str, int | None, list[str] | None]:
    # 轻量兜底：优先尝试 pypdf；若环境未安装则返回空文本。
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return "", None, None
    try:
        reader = PdfReader(str(path))
        texts: list[str] = []
        for page in reader.pages:
            extracted = (page.extract_text() or "").strip()
            texts.append(extracted)
        clean_pages = [x for x in texts if x]
        return "\n\n".join(clean_pages), len(reader.pages), texts
    except Exception:
        return "", None, None


def _read_image_file(path: Path) -> tuple[str, str]:
    try:
        from PIL import Image  # type: ignore
        import pytesseract  # type: ignore
    except Exception:
        return "", "ocr_python_deps_missing"
    try:
        with Image.open(path) as image:
            try:
                text = (pytesseract.image_to_string(image, lang="chi_sim+eng") or "").strip()
                if text:
                    return text, "ok"
            except Exception:
                pass
            try:
                text = (pytesseract.image_to_string(image, lang="eng") or "").strip()
                if text:
                    return text, "ocr_english_only"
            except Exception:
                pass
            return "", "ocr_binary_missing_or_failed"
    except Exception:
        return "", "ocr_binary_missing_or_failed"


def _ocr_image(image) -> tuple[str, str]:
    try:
        import pytesseract  # type: ignore
    except Exception:
        return "", "ocr_python_deps_missing"
    if not shutil.which("tesseract"):
        return "", "ocr_binary_missing_or_failed"
    try:
        try:
            text = (pytesseract.image_to_string(image, lang="chi_sim+eng") or "").strip()
            if text:
                return text, "ok"
        except Exception:
            pass
        try:
            text = (pytesseract.image_to_string(image, lang="eng") or "").strip()
            if text:
                return text, "ocr_english_only"
        except Exception:
            pass
        return "", "ocr_binary_missing_or_failed"
    except Exception:
        return "", "ocr_binary_missing_or_failed"


def _ocr_pdf_file(path: Path) -> tuple[str, int | None, list[str] | None, str]:
    try:
        import pypdfium2 as pdfium  # type: ignore
    except Exception:
        return "", None, None, "pdf_ocr_renderer_missing"
    try:
        pdf = pdfium.PdfDocument(str(path))
        page_count = len(pdf)
        page_texts: list[str] = []
        reasons: list[str] = []
        for index in range(page_count):
            page = pdf[index]
            bitmap = page.render(scale=2)
            image = bitmap.to_pil()
            text, reason = _ocr_image(image)
            page_texts.append(text)
            reasons.append(reason)
        combined = "\n\n".join(x for x in page_texts if x).strip()
        if combined:
            if any(reason == "ocr_english_only" for reason in reasons):
                return combined, page_count, page_texts, "ocr_english_only"
            return combined, page_count, page_texts, "ok"
        fallback_reason = next((reason for reason in reasons if reason != "ok"), "pdf_ocr_failed")
        return "", page_count, page_texts, fallback_reason
    except Exception:
        return "", None, None, "pdf_ocr_failed"


def _parse_hint_json(*, kind: str, ocr_suggested: bool, reason: str) -> str:
    return json.dumps(
        {
            "kind": kind,
            "ocr_suggested": ocr_suggested,
            "reason": reason,
        },
        ensure_ascii=False,
    )


def _merge_outline_json(base: str | None, **extra: object) -> str | None:
    payload: dict[str, object] = {}
    if base:
        try:
            payload.update(json.loads(base))
        except Exception:
            payload["raw_outline"] = base
    payload.update({k: v for k, v in extra.items() if v is not None})
    return json.dumps(payload, ensure_ascii=False) if payload else None


def _read_with_unstructured(path: Path) -> ParsedDocument | None:
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            from unstructured.partition.docx import partition_docx  # type: ignore

            elements = partition_docx(filename=str(path))
        elif suffix == ".pdf":
            from unstructured.partition.pdf import partition_pdf  # type: ignore

            elements = partition_pdf(
                filename=str(path),
                strategy="fast",
                infer_table_structure=False,
            )
        elif suffix in {".txt", ".md", ".markdown", ".html", ".htm"}:
            from unstructured.partition.auto import partition  # type: ignore

            elements = partition(filename=str(path))
        else:
            return None
    except Exception:
        return None

    lines: list[str] = []
    categories: dict[str, int] = {}
    for element in elements:
        text = str(getattr(element, "text", "") or "").strip()
        if not text:
            continue
        category = str(getattr(element, "category", "") or "").strip() or "Unknown"
        categories[category] = categories.get(category, 0) + 1
        if category in {"Title", "Header"}:
            if lines and lines[-1] != "":
                lines.append("")
            lines.append(text)
            lines.append("")
        else:
            lines.append(text)

    merged_text = "\n".join(lines).strip()
    if not merged_text:
        return None
    outline = _merge_outline_json(
        None,
        kind=suffix.lstrip(".") or "text",
        parser_engine="unstructured",
        parser_mode="element_partition",
        element_categories=categories,
    )
    return ParsedDocument(
        text=merged_text,
        page_count=None,
        outline_json=outline,
        pages=None,
        parser_engine="unstructured",
        parser_mode="element_partition",
    )


def parse_path(path: Path, preferred_engine: str = "auto") -> ParsedDocument:
    suffix = path.suffix.lower()
    if preferred_engine in {"auto", "unstructured"} and suffix in {".docx", ".txt", ".md", ".markdown", ".html", ".htm", ".pdf"}:
        parsed = _read_with_unstructured(path)
        if parsed:
            # PDF 优先保留当前内置按页/OCR 链路；只有 docx / text 类优先走结构化引擎。
            if suffix != ".pdf":
                return parsed
    if suffix in TEXT_EXTENSIONS:
        text = _read_text_file(path).strip()
        return ParsedDocument(text=text, parser_engine="builtin", parser_mode="plain_text")
    if suffix == ".docx":
        text = _read_docx_file(path).strip()
        return ParsedDocument(text=text, parser_engine="builtin", parser_mode="docx_paragraphs")
    if suffix == ".pdf":
        text, page_count, pages = _read_pdf_file(path)
        outline_json = None
        if page_count and not text.strip():
            ocr_text, ocr_page_count, ocr_pages, ocr_reason = _ocr_pdf_file(path)
            if ocr_text.strip():
                return ParsedDocument(
                    text=ocr_text.strip(),
                    page_count=ocr_page_count or page_count,
                    pages=ocr_pages,
                    outline_json=_merge_outline_json(
                        _parse_hint_json(
                            kind="pdf",
                            ocr_suggested=False,
                            reason=ocr_reason,
                        ),
                        parser_engine="builtin",
                        parser_mode="pdf_ocr_pages",
                    ),
                    parser_engine="builtin",
                    parser_mode="pdf_ocr_pages",
                )
            outline_json = _parse_hint_json(
                kind="pdf",
                ocr_suggested=True,
                reason=ocr_reason if ocr_reason != "ok" else "pdf_text_empty_maybe_scanned",
            )
        return ParsedDocument(
            text=text.strip(),
            page_count=page_count,
            pages=pages,
            outline_json=_merge_outline_json(
                outline_json,
                parser_engine="builtin",
                parser_mode="pdf_text_pages",
            ),
            parser_engine="builtin",
            parser_mode="pdf_text_pages",
        )
    if suffix in IMAGE_EXTENSIONS:
        text, image_reason = _read_image_file(path)
        text = text.strip()
        outline_json = None
        if not text:
            outline_json = _parse_hint_json(
                kind="image",
                ocr_suggested=True,
                reason=image_reason if image_reason != "ok" else "image_ocr_unavailable_or_no_text",
            )
        return ParsedDocument(
            text=text,
            page_count=1,
            outline_json=_merge_outline_json(
                outline_json,
                parser_engine="builtin",
                parser_mode="image_ocr",
            ),
            pages=[text] if text else None,
            parser_engine="builtin",
            parser_mode="image_ocr",
        )
    return ParsedDocument(text="", parser_engine="builtin", parser_mode="unsupported")
