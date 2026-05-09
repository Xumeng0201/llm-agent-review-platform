from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from docx import Document
from docx.shared import Pt

from .schemas import OverallReport


def build_report_docx_bytes(report: OverallReport) -> bytes:
    """生成与 HTML 版结构一致的 Word 文档（.docx）。"""
    doc = Document()
    normal = doc.styles["Normal"]
    if normal.font is not None:
        normal.font.size = Pt(11)

    doc.add_heading(report.task_name, level=0)
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
    doc.add_paragraph(f"方案评审报告 · 生成时间（UTC）：{ts}")

    doc.add_heading("评分汇总", level=1)
    doc.add_paragraph(
        f"总分：{report.total_score} / {report.max_total}"
        f"（共 {len(report.indicators)} 个一级指标，每项满分 10 分）"
    )

    doc.add_heading("综合结论", level=1)
    doc.add_paragraph().add_run(report.conclusion_label).bold = True
    for r in report.reasons:
        doc.add_paragraph(r, style="List Bullet")

    rs = (report.review_summary or "").strip()
    if rs:
        doc.add_heading("综合评审意见（智能体）", level=1)
        doc.add_paragraph(rs)

    doc.add_heading("分项得分", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    heads = table.rows[0].cells
    for i, h in enumerate(["#", "一级指标", "得分", "满分", "评审说明"]):
        heads[i].text = h
    for row in report.indicators:
        s = "—" if row.score is None else str(row.score)
        note = row.notes or "—"
        cells = table.add_row().cells
        cells[0].text = str(row.indicator_id)
        cells[1].text = row.title
        cells[2].text = s
        cells[3].text = str(row.max_score)
        cells[4].text = note

    doc.add_heading("项目亮点", level=1)
    doc.add_paragraph(report.summary_highlights or "（未填写）")

    doc.add_heading("存在问题与整改建议", level=1)
    doc.add_paragraph(report.summary_issues or "（未填写）")

    doc.add_heading("说明", level=1)
    doc.add_paragraph(
        "本报告依据系统当前配置的项目方案审核评测框架生成；"
        "可在 Microsoft Word 或 WPS 中打开本文件，并另存为 PDF。"
    )

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
